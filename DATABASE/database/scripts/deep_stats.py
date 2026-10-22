"""Deeper stats: categoricals, joins, SEARCH used cells, amounts."""
from __future__ import annotations

import json
import os
from collections import Counter

import pandas as pd
from openpyxl import load_workbook

PATH = r"c:\Users\vedan\Downloads\UG - Copy.xlsx"
OUT = r"C:\database\docs"


def main():
    data = pd.read_excel(PATH, sheet_name="DATA")
    data.columns = [str(c).strip() for c in data.columns]
    fee = pd.read_excel(PATH, sheet_name="FEE MAST")

    cats = [
        "Branch",
        "Eligible Category",
        "Alloted Category",
        "Domicile",
        "Gender",
        "EWS",
        "Status",
        "Round",
        "Payment Mode",
        "Religion",
        "Admission Year",
        "Remarks",
        "fee stru",
        "Ofline Admission Category",
    ]
    cat_out = {}
    for c in cats:
        if c not in data.columns:
            cat_out[c] = "MISSING"
            continue
        vc = data[c].fillna("(null)").astype(str).str.strip().value_counts()
        cat_out[c] = {k: int(v) for k, v in vc.items()}

    fill = {}
    for c in data.columns:
        nonempty = data[c].notna() & (data[c].astype(str).str.strip() != "") & (data[c].astype(str) != "nan")
        fill[c] = {"nonempty": int(nonempty.sum()), "unique": int(data.loc[nonempty, c].astype(str).nunique())}

    # duplicates Unique ID
    uid = data["Unique ID"]
    dups = data[uid.notna() & uid.duplicated(keep=False)][["Sr. No.", "Unique ID", "Enrolment No.", "Name", "Admission Year", "Branch"]].head(30)

    # fee join
    fee_keys = set(fee["FEE MAST"].dropna().astype(str).str.strip())
    data_fs = set(data["fee stru"].dropna().astype(str).str.strip())
    join = {
        "fee_mast_populated_rows": int(fee["FEE MAST"].notna().sum()),
        "fee_keys": sorted(fee_keys),
        "data_fee_stru_nunique": len(data_fs),
        "in_data_not_mast": sorted(data_fs - fee_keys)[:80],
        "in_mast_not_data": sorted(fee_keys - data_fs)[:80],
        "matched": sorted(data_fs & fee_keys),
    }

    numeric_cols = ["Total Amount", "Student submitted Amount", "DTE Fee"] + [
        c for c in data.columns if c.startswith("JUL-") or c.startswith("JAN-")
    ]
    sums = {}
    for c in numeric_cols:
        s = pd.to_numeric(data[c], errors="coerce")
        sums[c] = {
            "sum": float(s.fillna(0).sum()),
            "nonempty": int(s.notna().sum()),
            "min": None if s.dropna().empty else float(s.min()),
            "max": None if s.dropna().empty else float(s.max()),
        }

    # year & sem table on fee sheet
    fee2 = fee.copy()
    fee2.columns = [str(c) for c in fee2.columns]
    year_sem = fee2.iloc[:, [20, 21, 22]].dropna(how="all")
    year_sem.columns = ["year", "sem_code", "year_sem"]

    out = {
        "categoricals": cat_out,
        "fill": fill,
        "unique_id_dups_sample": dups.to_dict(orient="records"),
        "fee_join": join,
        "sums": sums,
        "year_sem_rows": int(len(year_sem)),
        "year_sem_unique": int(year_sem["year_sem"].nunique()),
        "sem_codes": sorted(year_sem["sem_code"].dropna().astype(str).unique().tolist()),
    }
    with open(os.path.join(OUT, "deep_stats.json"), "w", encoding="utf-8") as f:
        json.dump(out, f, indent=2, default=str)
    print("wrote deep_stats")
    print("enrolment unique", data["Enrolment No."].nunique(), "rows", len(data))
    print("aadhaar unique/nonempty", data["Adharcard No."].nunique(dropna=True), data["Adharcard No."].notna().sum())
    print("class roll unique", data["Class Roll No."].nunique())

    # SEARCH used cells via read_only
    print("scanning SEARCH read_only for any values...")
    wb = load_workbook(PATH, read_only=True, data_only=True, keep_vba=False)
    ws = wb["SEARCH"]
    nonempty = []
    formulas_or_text = Counter()
    n = 0
    for row in ws.iter_rows(max_row=200, max_col=80, values_only=False):
        for cell in row:
            n += 1
            if cell.value is not None and cell.value != "":
                nonempty.append((cell.coordinate, str(cell.value)[:200]))
            if len(nonempty) >= 80:
                break
        if len(nonempty) >= 80:
            break
    print("checked cells", n, "found", len(nonempty))
    print(nonempty[:40])
    wb.close()

    with open(os.path.join(OUT, "search_cells.json"), "w", encoding="utf-8") as f:
        json.dump({"found": nonempty}, f, indent=2)

    # DATA first row sample as dict
    print(data.iloc[0].to_dict())


if __name__ == "__main__":
    main()
